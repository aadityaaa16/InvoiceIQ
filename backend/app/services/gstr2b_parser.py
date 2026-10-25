"""Parse GSTR-2B data (what suppliers reported to the GST portal) into flat records.

Supported inputs:
- The GST portal's JSON download: {"data": {"docdata": {"b2b": [{"ctin", "trdnm", "inv": [...]}]}}}
- A simple Excel sheet with one row per invoice and these headers (any order, any case):
  GSTIN | Supplier Name | Invoice Number | Invoice Date | Invoice Value |
  Taxable Value | IGST | CGST | SGST
"""

import io
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

from app.core.exceptions import BadRequestError
from app.utils.gstin import normalize_gstin


class GstrRecord(BaseModel):
    supplier_gstin: str
    supplier_name: str | None = None
    invoice_number: str
    invoice_date: date | None = None
    invoice_value: Decimal  # total including tax
    taxable_value: Decimal | None = None
    igst: Decimal = Decimal("0")
    cgst: Decimal = Decimal("0")
    sgst: Decimal = Decimal("0")

    @property
    def total_tax(self) -> Decimal:
        return self.igst + self.cgst + self.sgst


def _parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):  # portal uses DD-MM-YYYY
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognised date '{value}'")


def _money(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        # str() first: Excel gives floats, and Decimal(float) would keep binary noise.
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f"Invalid amount '{value}'") from exc


def _record(raw: dict[str, Any], row: str) -> GstrRecord:
    try:
        return GstrRecord(
            supplier_gstin=normalize_gstin(str(raw["gstin"])),
            supplier_name=raw.get("name"),
            invoice_number=str(raw["inum"]).strip(),
            invoice_date=_parse_date(raw.get("dt")),
            invoice_value=_money(raw.get("val")),
            taxable_value=_money(raw.get("txval")) if raw.get("txval") is not None else None,
            igst=_money(raw.get("igst")),
            cgst=_money(raw.get("cgst")),
            sgst=_money(raw.get("sgst")),
        )
    except (KeyError, ValueError, ValidationError) as exc:
        raise BadRequestError(f"Invalid GSTR-2B record at {row}: {exc}") from exc


def parse_json(data: bytes) -> list[GstrRecord]:
    try:
        document = json.loads(data)
    except json.JSONDecodeError as exc:
        raise BadRequestError("File is not valid JSON") from exc
    container = (
        document.get("data", {}).get("docdata", document) if isinstance(document, dict) else {}
    )
    suppliers = container.get("b2b") if isinstance(container, dict) else None
    if not isinstance(suppliers, list):
        raise BadRequestError("GSTR-2B JSON must contain a 'b2b' list")

    records = []
    for s_index, supplier in enumerate(suppliers):
        for i_index, invoice in enumerate(supplier.get("inv", [])):
            raw = {"gstin": supplier.get("ctin"), "name": supplier.get("trdnm"), **invoice}
            records.append(_record(raw, f"b2b[{s_index}].inv[{i_index}]"))
    return records


_EXCEL_HEADERS = {
    "gstin": "gstin",
    "gstin of supplier": "gstin",
    "supplier gstin": "gstin",
    "supplier name": "name",
    "trade name": "name",
    "invoice number": "inum",
    "invoice date": "dt",
    "invoice value": "val",
    "taxable value": "txval",
    "igst": "igst",
    "cgst": "cgst",
    "sgst": "sgst",
}


def parse_excel(data: bytes) -> list[GstrRecord]:
    try:
        # read_only + data_only: stream rows and read cached formula results, not formulas.
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
    except Exception as exc:  # openpyxl raises several unrelated types for bad files
        raise BadRequestError("File is not a valid Excel workbook") from exc

    rows = sheet.iter_rows(values_only=True)
    header = [str(cell or "").strip().lower() for cell in next(rows, ())]
    columns = {
        index: _EXCEL_HEADERS[name] for index, name in enumerate(header) if name in _EXCEL_HEADERS
    }
    missing = {"gstin", "inum", "val"} - set(columns.values())
    if missing:
        raise BadRequestError(
            "Excel sheet is missing required columns",
            details={"required": ["GSTIN", "Invoice Number", "Invoice Value"]},
        )

    records = []
    for row_number, row in enumerate(rows, start=2):
        if not any(row):
            continue  # skip blank lines
        raw = {key: row[index] for index, key in columns.items() if index < len(row)}
        records.append(_record(raw, f"row {row_number}"))
    return records


def parse_gstr2b(filename: str, data: bytes) -> list[GstrRecord]:
    if filename.lower().endswith(".json"):
        return parse_json(data)
    if filename.lower().endswith(".xlsx"):
        return parse_excel(data)
    raise BadRequestError("GSTR-2B file must be .json or .xlsx")
